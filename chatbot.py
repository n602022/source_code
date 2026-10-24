# -*- coding: utf-8 -*-
import json
import os
import re
import sqlite3
import sys
import unicodedata
from pathlib import Path
from docx import Document

BASE_DIR = Path(__file__).resolve().parent
THU_MUC_TAI_LIEU = BASE_DIR / "safety_docs"
DB_PATH = BASE_DIR / "safeguard.db"
EVIDENCE_DIR = BASE_DIR / "evidence"

TEN_VAN_BAN_DEP = {
    "Quy_dinh_trang_thiet_bi_bao_ho.txt": "Điều 1 — Quy định Trang thiết bị Bảo hộ Cá nhân (PPE)",
    "Quy_dinh_vung_nguy_hiem_DangerZone.txt": "Điều 2 — Quy định Vùng Nguy hiểm & Dừng Khẩn cấp",
    "Quyen_va_nghia_vu_nguoi_lao_dong.txt": "Điều 3 — Quyền và Nghĩa vụ An toàn của Người Lao động",
    "Xu_ly_vi_pham_qua_camera_AI.txt": "Điều 4 — Quy trình Xử lý Vi phạm qua Hệ thống Camera AI",
    "An_toan_han_cat_va_lam_viec_tren_cao.txt": "Điều 5 — An toàn Làm việc trên Cao & Hàn cắt Kim loại",
}

GIAI_THICH_LOI_VI_PHAM = {
    "NO-Hardhat": {
        "ten_tv": "Không đội mũ bảo hộ",
        "muc_do": "Nguy cơ cao (Chấn thương vùng đầu khi có vật rơi hoặc va đập)",
        "cach_duyet": "Chọn Xác nhận vi phạm nếu công nhân để đầu trần hoặc chỉ đội mũ vải, mũ lưỡi trai. Chọn Đánh dấu báo nhầm nếu công nhân có đội mũ cứng bảo hộ nhưng bị bóng râm hoặc góc máy che khuất.",
        "xu_ly": "Phát loa nhắc nhở trực tiếp tại hiện trường (lần 1); lập biên bản trừ lương hoặc kỷ luật (từ lần 2).",
        "nguon": "Điều 1 & Điều 4 — Nội quy An toàn",
    },
    "NO-Safety Vest": {
        "ten_tv": "Không mặc áo phản quang",
        "muc_do": "Nguy cơ cao (Khó nhận diện khi có xe nâng hoặc máy móc di chuyển)",
        "cach_duyet": "Chọn Xác nhận vi phạm nếu công nhân mặc áo thường không có dải phản quang. Chọn Đánh dấu báo nhầm nếu áo phản quang bị lóa sáng do tia lửa hàn hoặc bị khuất ở rìa khung hình.",
        "xu_ly": "Yêu cầu bổ sung áo phản quang ngay lập tức; tái phạm lập biên bản kỷ luật.",
        "nguon": "Điều 1 & Điều 4 — Nội quy An toàn",
    },
    "NO-Mask": {
        "ten_tv": "Không đeo khẩu trang bảo hộ",
        "muc_do": "Nguy cơ trung bình (Rủi ro hít khói bụi, khí hàn và mạt kim loại)",
        "cach_duyet": "Chỉ chọn Xác nhận vi phạm khi nhìn rõ chính diện khuôn mặt không đeo khẩu trang. Nếu công nhân đang quay lưng hoặc cúi mặt, chọn Đánh dấu báo nhầm.",
        "xu_ly": "Nhắc nhở đeo khẩu trang chuyên dụng tại khu vực nhiều bụi và xưởng hàn cắt.",
        "nguon": "Điều 1 & Điều 4 — Nội quy An toàn",
    },
}

TRI_THUC_GIAM_SAT = [
    {
        "tu_khoa": ("bam nham", "duyet nham", "sua lai", "hoan tac", "undo", "doi trang thai", "sai thanh dung", "dung thanh sai"),
        "tra_loi": (
            "**Hướng dẫn sửa lại khi bấm duyệt nhầm (Đúng ↔ Sai):**\n"
            "- **Cách 1 (Đổi nhanh ngay sau khi duyệt):** Bấm nút **Đổi sang Báo nhầm** (hoặc **Đổi sang Vi phạm thật**) trên thanh thông báo hoàn tác ở đầu **Tab 2 (Duyệt Ảnh Vi Phạm)**.\n"
            "- **Cách 2 (Sửa trong danh sách đã duyệt):** Kéo xuống bảng **Lịch sử sự kiện đã duyệt & Chỉnh sửa trạng thái** ở cuối Tab 2, chọn mã sự kiện và bấm chuyển đổi trực tiếp giữa **Vi phạm thật**, **Báo nhầm**, hoặc **Chờ duyệt**."
        ),
    },
    {
        "tu_khoa": ("nhieu lan", "chup lien tuc", "bi nhieu", "1 nguoi nhieu loi", "trung lap", "khoa nguoi", "tracking"),
        "tra_loi": (
            "**Cơ chế Khóa Công Nhân Chống Báo Lặp & Tự Động Lọc Ánh Sáng Lóe:**\n"
            "- **Không chụp lặp liên tục:** Hệ thống định danh từng người (`CN #01`, `CN #02`...), chỉ lấy lỗi chính rõ nhất và **khóa công nhân đã ghi nhận vi phạm** để không chụp ảnh lặp lại qua từng giây.\n"
            "- **Thích ứng ánh sáng lóe & độ mờ:** Khi phát hiện tia sáng chói (như tia lửa hàn) hoặc khung hình rung mờ, AI tự động nâng ngưỡng kiểm tra để không bắt nhầm công nhân đang làm việc thành vi phạm."
        ),
    },
    {
        "tu_khoa": ("anh sang", "loe sang", "bi mo", "nguoc sang", "choi sang", "tia lua han", "hoan canh", "am thanh"),
        "tra_loi": (
            "**Cơ chế Tự Động Thích Ứng Ánh Sáng Lóe & Độ Mờ:**\n"
            "- Hệ thống tự động đo tỷ lệ điểm ảnh chói lóa (như tia lửa hàn) và độ mờ rung của từng khung hình.\n"
            "- Khi phát hiện ánh sáng lóe mạnh làm trắng xóa vùng áo/mũ hoặc khung hình bị mờ, AI tự động lọc bỏ các cảnh báo độ tin cậy thấp để **tránh bắt nhầm người đang làm việc thành vi phạm**."
        ),
    },
    {
        "tu_khoa": ("xuat file", "tai bao cao", "excel", "csv", "luu bang chung", "bien ban", "nop bao cao", "zip"),
        "tra_loi": (
            "**Hướng dẫn Xuất File Báo Cáo Ca Làm & Lưu Bằng Chứng:**\n"
            "- Tại cột bên phải của **Tab 1 (Quét AI & Thống Kê Vi Phạm)**, mục **3. Xuất hồ sơ vi phạm ca làm & bằng chứng**, bạn bấm **Xuất file báo cáo (.CSV / Excel)** hoặc **Tải gói ảnh bằng chứng (.ZIP)**.\n"
            "- Hồ sơ tải về bao gồm bảng tổng hợp đầy đủ thời gian, khu vực camera, mã công nhân, hành vi vi phạm, trạng thái duyệt và toàn bộ ảnh chụp hiện trường."
        ),
    },
    {
        "tu_khoa": ("do nhay", "nguong phat hien", "toi thieu", "tu dong xac nhan", "thanh truot", "cau hinh"),
        "tra_loi": (
            "**Ý nghĩa 2 thanh trượt cấu hình AI ở bảng điều khiển bên trái:**\n"
            "- **Độ nhạy phát hiện (tối thiểu):** Mức độ chắc chắn tối thiểu để AI bắt đầu khoanh vùng sự kiện (mặc định `0.50`). Mọi cảnh báo thấp hơn mức này sẽ bị bỏ qua ngay.\n"
            "- **Ngưỡng AI tự động xác nhận lỗi:** Nếu độ tin cậy của AI đạt từ ngưỡng này trở lên (mặc định `0.75`), hệ thống tự động xác nhận vi phạm; nếu nằm giữa *Độ nhạy tối thiểu* và *Ngưỡng tự động xác nhận*, hệ thống sẽ chụp ảnh gửi sang **Tab 2** để Giám sát viên duyệt."
        ),
    },
    {
        "tu_khoa": ("nong qua", "oi buc", "coi ao", "bo mu", "nghi ngoi", "tam thoi", "mot luc"),
        "tra_loi": (
            "**Xử lý trường hợp công nhân tự ý tháo mũ/áo vì thời tiết nóng:**\n"
            "- **Quyết định:** Không chấp nhận việc tháo mũ bảo hộ hoặc áo phản quang khi đang đứng trong khu vực thi công, vận hành máy.\n"
            "- **Hướng xử lý:** Bấm **Xác nhận vi phạm**, sau đó nhắc công nhân di chuyển ra **khu vực nghỉ giải lao quy định** nếu muốn tháo đồ bảo hộ *(Căn cứ: Điều 1 & Điều 3)*."
        ),
    },
    {
        "tu_khoa": ("khach", "tham quan", "doi tac", "ky su", "quan ly vao", "nguoi ngoai"),
        "tra_loi": (
            "**Quy định đối với Khách tham quan / Kỹ sư kiểm tra xưởng:**\n"
            "- Tất cả cá nhân bước vào khu vực sản xuất đều bắt buộc đội mũ bảo hộ và mặc áo phản quang, không có ngoại lệ.\n"
            "- Tuyệt đối không để người không phận sự đi vào vùng máy móc nguy hiểm đang vận hành *(Căn cứ: Điều 1 & Điều 2)*."
        ),
    },
    {
        "tu_khoa": ("tai nan", "bi thuong", "bong", "dien giat", "nga", "cap cuu", "so cuu", "chay no"),
        "tra_loi": (
            "**Quy trình 3 bước ứng phó khẩn cấp khi xảy ra sự cố/tai nạn:**\n"
            "1. **Dừng máy khẩn cấp:** Ngắt nguồn điện hoặc bấm nút dừng khẩn cấp dây chuyền tại khu vực xảy ra sự cố.\n"
            "2. **Cấp cứu tại chỗ:** Đưa nạn nhân ra khỏi vùng nguy hiểm, sơ cứu và gọi bộ phận Y tế.\n"
            "3. **Trích xuất bằng chứng:** Tải gói ảnh bằng chứng và báo cáo CSV trên hệ thống để lập biên bản điều tra *(Căn cứ: Điều 2 & Điều 4)*."
        ),
    },
    {
        "tu_khoa": ("tai pham", "lan 2", "lan 3", "co tinh", "chong doi"),
        "tra_loi": (
            "**Chế tài xử lý công nhân cố tình tái phạm (Điều 4):**\n"
            "- **Vi phạm lần 1:** Hệ thống phát tín hiệu âm thanh và nhắc nhở trực tiếp qua loa tại khu vực làm việc.\n"
            "- **Tái phạm từ lần 2 trở đi:** Bấm **Xác nhận vi phạm**, xuất file báo cáo kèm ảnh chụp bằng chứng để áp dụng hình thức **trừ lương hoặc kỷ luật**, đồng thời tạm dừng công việc cho đến khi trang bị đủ đồ bảo hộ."
        ),
    },
    {
        "tu_khoa": ("het do", "kho het", "bi hong", "bi nut", "rach", "thay moi", "cap phat"),
        "tra_loi": (
            "**Xử lý khi thiết bị bảo hộ bị nứt hỏng, nứt vỡ hoặc hư hỏng (Điều 3):**\n"
            "- Không cho phép người lao động tiếp tục làm việc với mũ bảo hộ bị nứt hoặc dây đai an toàn bị sờn đứt.\n"
            "- Yêu cầu công nhân báo ngay cho bộ phận Kho để cấp phát thiết bị mới; trong thời gian chờ đợi phải đứng ngoài khu vực máy móc."
        ),
    },
    {
        "tu_khoa": ("mu luoi trai", "non ket", "non vai", "mu vai", "non thuong"),
        "tra_loi": (
            "**Quy chuẩn Mũ bảo hộ hợp lệ (Điều 1 & Điều 5):**\n"
            "- Mũ lưỡi trai hoặc mũ vải thông thường **không có khả năng chống va đập và vật rơi**, do đó vẫn tính là vi phạm **Không đội mũ bảo hộ**.\n"
            "- Giám sát viên cần bấm **Xác nhận vi phạm** và yêu cầu công nhân đổi sang mũ bảo hộ vỏ cứng có cài quai."
        ),
    },
    {
        "tu_khoa": ("tai sao can duyet", "khi nao can duyet", "ai tu duyet", "cho nguoi duyet", "can nguoi", "khong chac chan", "nghi ngo"),
        "tra_loi": (
            "**Cơ chế Phối hợp giữa AI Tự Động và Giám Sát Viên:**\n"
            "- **Khi AI chắc chắn cao (đạt ngưỡng tự động xác nhận):** Hệ thống tự động ghi nhận vào thống kê ca trực để xử lý tức thì.\n"
            "- **Khi AI nghi ngờ (nằm giữa độ nhạy tối thiểu và ngưỡng tự động xác nhận):** Hệ thống tự động chụp màn hình và chuyển sang **Tab 2** để Giám sát viên kiểm tra bằng mắt thường."
        ),
    },
    {
        "tu_khoa": ("quyen rieng tu", "ap luc", "theo doi", "cong nhan phan doi", "giam sat 24", "dao duc", "bao mat"),
        "tra_loi": (
            "**Chính sách Bảo mật & Quyền riêng tư của Người lao động:**\n"
            "- Hệ thống chỉ nhận diện trang bị bảo hộ (mũ, áo phản quang, khẩu trang) để bảo vệ tính mạng người lao động, không thu thập sinh trắc học khuôn mặt cá nhân.\n"
            "- Toàn bộ video được xử lý nội bộ tại doanh nghiệp và chỉ lưu lại ảnh chụp khi phát hiện dấu hiệu vi phạm."
        ),
    },
    {
        "tu_khoa": ("han cat", "han xi", "tho han", "kinh han"),
        "tra_loi": (
            "**Quy định An toàn Hàn cắt Kim loại (Điều 5):**\n"
            "- Thợ hàn bắt buộc trang bị đầy đủ: **mặt nạ hàn chuyên dụng, kính bảo hộ, găng tay da chịu nhiệt và quần áo bảo hộ dài tay chống cháy**.\n"
            "- Nếu thiếu bất kỳ thiết bị nào, yêu cầu tạm dừng thao tác hàn ngay để phòng ngừa bỏng tia lửa và tổn thương mắt."
        ),
    },
    {
        "tu_khoa": ("tren cao", "gian giao", "leo cao", "day dai"),
        "tra_loi": (
            "**Quy định Làm việc trên Cao & Giàn giáo (Điều 5):**\n"
            "- Công nhân làm việc trên giàn giáo hoặc trên cao bắt buộc phải đeo **dây đai an toàn** móc cố định và đội **mũ bảo hộ cài quai chắc chắn**.\n"
            "- Giám sát viên phải đình chỉ công việc ngay nếu phát hiện công nhân leo cao không sử dụng dây đai an toàn."
        ),
    },
    {
        "tu_khoa": ("vung nguy hiem", "danger zone", "vung cam", "khu vuc cam", "dung khan cap", "may moc"),
        "tra_loi": (
            "**Quy định Kiểm soát Vùng Nguy hiểm (Điều 2):**\n"
            "- Khu vực máy móc đang vận hành có rào chắn hoặc vạch giới hạn đỏ là vùng nguy hiểm, cấm người không phận sự tự ý bước vào.\n"
            "- Khi phát hiện người xâm nhập trái phép vào vùng nguy hiểm, hệ thống phát cảnh báo và kích hoạt quy trình dừng khẩn cấp dây chuyền."
        ),
    },
    {
        "tu_khoa": ("tu choi", "nguy co", "quyen", "nghia vu"),
        "tra_loi": (
            "**Quyền và Nghĩa vụ An toàn của Người lao động (Điều 3):**\n"
            "- Người lao động có quyền **từ chối làm việc hoặc rời khỏi nơi làm việc** khi thấy rõ nguy cơ xảy ra tai nạn lao động đe dọa tính mạng, sức khỏe.\n"
            "- Có nghĩa vụ tuân thủ quy định bảo hộ và báo cáo kịp thời khi phát hiện thiết bị mất an toàn."
        ),
    },
    {
        "tu_khoa": ("xu ly", "xu phat", "ky luat", "tru luong", "loa nhac nho", "quy trinh", "gui loa", "dung khan cap day chuyen"),
        "tra_loi": (
            "**Quy trình Xử lý Vi phạm & Điều khiển Khẩn cấp IoT (Điều 2 & Điều 4):**\n"
            "- **Gửi loa nhắc nhở (Vi phạm lần 1):** Đội trưởng bấm nút **Gửi loa nhắc nhở** tại Tab 2 để phát âm báo trực tiếp qua loa truyền thanh tại phân xưởng.\n"
            "- **Dừng khẩn cấp dây chuyền (Nguy cơ tai nạn cao):** Đội trưởng bấm nút **Dừng khẩn cấp dây chuyền** tại Tab 2 để ngắt tín hiệu máy móc tức thì và lưu biên bản xử lý."
        ),
    },
    {
        "tu_khoa": ("mu bao ho", "mu bao hiem", "non bao ho", "khong doi mu", "no-hardhat"),
        "tra_loi": (
            "**Quy định Đội mũ bảo hộ (Điều 1 & Điều 4):**\n"
            "- Tất cả công nhân làm việc tại công trường và phân xưởng bắt buộc phải đội mũ bảo hộ đạt chuẩn trong suốt ca làm việc để phòng chống vật rơi, va đập.\n"
            "- Vi phạm sẽ bị nhắc nhở qua loa (lần 1) hoặc lập biên bản kỷ luật (từ lần 2)."
        ),
    },
    {
        "tu_khoa": ("ao phan quang", "ao bao ho", "khong mac ao", "no-safety vest"),
        "tra_loi": (
            "**Quy định Mặc áo phản quang (Điều 1 & Điều 4):**\n"
            "- Người lao động bắt buộc mặc áo phản quang khi làm việc tại công trường, khu vực thiếu sáng hoặc nơi có phương tiện, máy móc di chuyển.\n"
            "- Mọi trường hợp thiếu áo phản quang đều được hệ thống lưu ảnh bằng chứng để nhắc nhở và xử lý theo nội quy."
        ),
    },
    {
        "tu_khoa": ("khau trang", "gang tay", "khoi bui", "no-mask"),
        "tra_loi": (
            "**Quy định Khẩu trang & Găng tay bảo hộ (Điều 1):**\n"
            "- Bắt buộc đeo khẩu trang tại khu vực phát sinh khói bụi, hóa chất và đeo găng tay bảo hộ khi vận hành máy móc hoặc tiếp xúc vật sắc nhọn."
        ),
    },
]

TU_DONG_NGHIA = {
    "no-hardhat": "khong doi mu bao ho",
    "no hardhat": "khong doi mu bao ho",
    "hardhat": "mu bao ho",
    "no-safety vest": "khong mac ao phan quang",
    "no safety vest": "khong mac ao phan quang",
    "safety vest": "ao phan quang",
    "no-mask": "khong deo khau trang",
    "no mask": "khong deo khau trang",
    "danger zone": "vung nguy hiem",
    "mu bao hiem": "mu bao ho",
    "non bao ho": "mu bao ho",
    "non bao hiem": "mu bao ho",
    "do bao ho": "thiet bi bao ho",
    "ppe": "bao ho ca nhan",
}

TU_BO_QUA = {
    "la", "va", "cua", "co", "the", "nao", "gi", "khi", "thi", "ma", "de",
    "cau", "lien",
    "o", "tai", "cho", "voi", "duoc", "nhung", "cac", "mot", "nhu", "ra",
    "cong", "nhan", "nguoi", "lao", "dong", "lam", "viec", "nha", "anh", "chi",
    "may", "truong", "bi", "phat", "bao", "nhieu", "tien", "ngay", "gio",
    "nam", "sao", "tinh", "khong", "phai", "quy", "dinh", "hoi", "biet",
    "neu", "thay", "tren", "trong", "ngoai", "can", "nen", "ai", "ve", "theo",
    "tui", "toi", "minh", "ban", "em", "ad", "bot", "oi", "nha", "nhe"
}


def chuan_hoa(text):
    text = str(text).lower().replace("đ", "d")
    text = unicodedata.normalize("NFD", text)
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    text = re.sub(r"[^a-z0-9\s\-]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def doc_file_docx(duong_dan):
    doc = Document(duong_dan)
    cac_dong = [p.text for p in doc.paragraphs if p.text.strip()]
    for bang in doc.tables:
        for hang in bang.rows:
            for o in hang.cells:
                if o.text.strip():
                    cac_dong.append(o.text)
    return "\n".join(cac_dong)


def tach_cau(van_ban):
    cac_cau = re.split(r"(?<=[.!?;])\s+|\n+", van_ban)
    return [c.strip() for c in cac_cau if c.strip()]


def nap_tai_lieu(thu_muc):
    if not os.path.isdir(thu_muc):
        return []
    danh_sach = []
    for ten_file in sorted(os.listdir(thu_muc)):
        if ten_file.startswith("~$"):
            continue
        duong_dan = os.path.join(thu_muc, ten_file)
        if ten_file.lower().endswith(".docx"):
            van_ban = doc_file_docx(duong_dan)
        elif ten_file.lower().endswith(".txt"):
            with open(duong_dan, "r", encoding="utf-8") as tep:
                van_ban = tep.read()
        else:
            continue
        khop = re.match(r"\s*(Điều\s*\d+)", van_ban)
        dieu = khop.group(1) if khop else "Nội quy An toàn"
        for cau in tach_cau(van_ban):
            danh_sach.append({"cau": cau, "nguon": ten_file, "dieu": dieu})
    return danh_sach


def lay_thong_ke_db():
    if not DB_PATH.is_file():
        return None
    try:
        with sqlite3.connect(DB_PATH) as conn:
            conn.row_factory = sqlite3.Row
            rows = conn.execute(
                "SELECT event_id, timestamp, violation_type, status FROM events ORDER BY event_id DESC"
            ).fetchall()
        total = len(rows)
        pending = sum(1 for r in rows if r["status"] == "pending")
        confirmed = sum(1 for r in rows if r["status"] == "confirmed")
        false_pos = sum(1 for r in rows if r["status"] == "false_positive")
        by_type = {}
        for r in rows:
            if r["status"] != "false_positive":
                vt = r["violation_type"]
                by_type[vt] = by_type.get(vt, 0) + 1
        latest = dict(rows[0]) if rows else None
        return {
            "total": total,
            "pending": pending,
            "confirmed": confirmed,
            "false_positive": false_pos,
            "by_type": by_type,
            "latest": latest,
        }
    except Exception:
        return None


def tim_kiem(tu_khoa, danh_sach, toi_da=2):
    q = chuan_hoa(tu_khoa)
    if not q:
        return []
    cac_cum = [q]
    for tu, thay_the in TU_DONG_NGHIA.items():
        if tu in q:
            cac_cum.append(q.replace(tu, thay_the))

    ket_qua = [m for m in danh_sach if any(c in chuan_hoa(m["cau"]) for c in cac_cum if len(c.split()) >= 2)]
    if ket_qua:
        return ket_qua[:toi_da]

    cac_tu = {t for c in cac_cum for t in c.split() if t not in TU_BO_QUA and len(t) > 1}
    if not cac_tu:
        return []

    co_diem = []
    for m in danh_sach:
        tu_trong_cau = set(chuan_hoa(m["cau"]).split())
        diem = len(cac_tu & tu_trong_cau)
        nguong = max(2, (len(cac_tu) + 1) // 2)
        if diem >= nguong:
            co_diem.append((diem, m))
    co_diem.sort(key=lambda x: -x[0])
    return [m for _, m in co_diem[:toi_da]]


def ask_chatbot(cau_hoi, selected_event=None):
    """Hàm trả lời nghiệp vụ chuẩn doanh nghiệp cho Giám sát viên."""
    q = chuan_hoa(cau_hoi)
    if not q:
        return "Vui lòng nhập câu hỏi hoặc chọn một chủ đề nghiệp vụ bên trên."

    words = set(q.split())

    if q in {"chao", "hi", "hello", "alo", "xin chao"} or (len(words) <= 3 and words & {"chao", "hi", "hello", "alo"}):
        stats = lay_thong_ke_db()
        cau_db = f" Hiện có **{stats['pending']} sự kiện đang chờ duyệt**." if stats and stats["pending"] > 0 else ""
        return (
            f"**Xin chào Giám sát viên.**{cau_db}\n"
            "Trợ lý AI sẵn sàng hỗ trợ: **(1) Hướng dẫn xử lý tình huống vi phạm tại hiện trường**, "
            "**(2) Cách sửa lại khi bấm duyệt nhầm hoặc xuất báo cáo ca làm**, và **(3) Tra cứu 5 Điều nội quy an toàn lao động**."
        )

    if any(k in q for k in ("cam on", "thank", "ok", "tot lam", "hay qua")):
        return "Đã ghi nhận. Nếu cần tra cứu thêm quy định hoặc hướng dẫn vận hành hệ thống, bạn hãy nhập câu hỏi."

    if selected_event and any(k in q for k in ("su kien nay", "loi nay", "anh nay", "duyet sao", "duyet the nao")):
        v_type = selected_event.get("violation_type", "NO-Hardhat")
        ev_id = selected_event.get("event_id", "?")
        info = GIAI_THICH_LOI_VI_PHAM.get(v_type, GIAI_THICH_LOI_VI_PHAM["NO-Hardhat"])
        return (
            f"**Hướng dẫn duyệt Mã #{ev_id} ({info['ten_tv']}):**\n"
            f"- **Cách kiểm tra ảnh:** {info['cach_duyet']}\n"
            f"- **Hình thức xử lý:** {info['xu_ly']} *({info['nguon']})*."
        )

    if any(
        k in q
        for k in (
            "thong ke",
            "bao nhieu loi",
            "bao nhieu vi pham",
            "tinh hinh ca",
            "tinh hinh vi pham",
            "so luong vi pham",
            "so luong loi",
            "chua duyet",
            "muc do an toan",
            "bao cao thong ke",
            "bao cao ca",
        )
    ) and not any(x in q for x in ("xuat file", "tai file", "excel", "csv", "zip", "luu bang chung")):
        stats = lay_thong_ke_db()
        if not stats or stats["total"] == 0:
            return "**Báo cáo ca làm việc:** Hiện tại chưa phát sinh cảnh báo vi phạm nào, tình trạng an toàn đạt 100%."
        chi_tiet = ", ".join(
            f"**{GIAI_THICH_LOI_VI_PHAM.get(k, {}).get('ten_tv', k)}**: {v}"
            for k, v in stats["by_type"].items()
        ) or "Không có vi phạm hợp lệ"
        return (
            f"**Tổng hợp tình hình an toàn ca làm việc:**\n"
            f"- **Tổng số cảnh báo đã ghi nhận:** **{stats['total']}** (Đã xác nhận vi phạm: **{stats['confirmed']}** | Đang chờ duyệt: **{stats['pending']}** | Đã loại bỏ do báo nhầm: **{stats['false_positive']}**).\n"
            f"- **Phân loại theo hành vi:** {chi_tiet}."
        )

    if any(k in q for k in ("tom tat", "toan bo", "5 dieu", "noi quy", "tat ca quy dinh")):
        return (
            "**Tóm tắt 5 Quy định An toàn Lao động cốt lõi:**\n"
            "1. **Điều 1 (Trang bị PPE):** Bắt buộc đội mũ bảo hộ, đeo găng tay, khẩu trang và mặc áo phản quang.\n"
            "2. **Điều 2 (Vùng nguy hiểm):** Cấm người không phận sự tự ý vào khu vực máy móc nguy hiểm; nếu xâm nhập trái phép phải dừng khẩn cấp.\n"
            "3. **Điều 3 (Quyền người lao động):** Được quyền từ chối làm việc khi có nguy cơ mất an toàn; thiết bị hư hỏng phải báo kho cấp mới.\n"
            "4. **Điều 4 (Chế tài xử lý):** Vi phạm lần 1 nhắc nhở qua loa; tái phạm từ lần 2 lập biên bản trừ lương hoặc kỷ luật.\n"
            "5. **Điều 5 (Trên cao & Hàn cắt):** Làm việc trên giàn giáo phải đeo dây đai an toàn + mũ cài quai; hàn cắt phải có mặt nạ hàn, kính, găng da và áo chống cháy."
        )

    for item in TRI_THUC_GIAM_SAT:
        if any(tk in q for tk in item["tu_khoa"]):
            return item["tra_loi"]

    danh_sach = nap_tai_lieu(THU_MUC_TAI_LIEU)
    ket_qua = tim_kiem(cau_hoi, danh_sach, toi_da=2)
    if not ket_qua:
        return "Chưa đủ tài liệu để trả lời"

    cac_dong = []
    for m in ket_qua:
        ten_dep = TEN_VAN_BAN_DEP.get(m["nguon"], m["dieu"])
        cac_dong.append(f"- {m['cau']} *(Căn cứ: {ten_dep})*")
    return "\n".join(cac_dong)


def main():
    while True:
        try:
            tu_khoa = input("\nCâu hỏi: ").strip()
        except (EOFError, KeyboardInterrupt):
            break
        if chuan_hoa(tu_khoa) in ("thoat", "exit", "quit"):
            break
        print("\nTrợ lý AI:\n" + ask_chatbot(tu_khoa))


if __name__ == "__main__":
    main()